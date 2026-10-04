"""Document exporter.

Generates downloadable PDF, DOCX, TXT and Markdown files
from conversations or AI responses.
"""
import io
import logging
from datetime import datetime
from typing import List, Dict, Optional
from pathlib import Path

logger = logging.getLogger(__name__)


class DocumentExporter:
    """Exports conversations or text into downloadable files."""
    
    SUPPORTED_FORMATS = {'txt', 'md', 'pdf', 'docx'}
    
    @classmethod
    def export(
        cls,
        content: str,
        format: str = 'txt',
        title: Optional[str] = None,
        metadata: Optional[Dict] = None
    ) -> bytes:
        """Export content to the given format.
        
        Args:
            content: Text content (markdown supported for PDF/DOCX)
            format: Target format (txt, md, pdf, docx)
            title: Document title
            metadata: Optional metadata dict
            
        Returns:
            Raw bytes of the generated document
        """
        format = format.lower().strip()
        
        if format not in cls.SUPPORTED_FORMATS:
            raise ValueError(
                f"Formato '{format}' no soportado. "
                f"Formatos permitidos: {', '.join(cls.SUPPORTED_FORMATS)}"
            )
        
        if format == 'txt':
            return cls._to_txt(content, title, metadata)
        elif format == 'md':
            return cls._to_markdown(content, title, metadata)
        elif format == 'pdf':
            return cls._to_pdf(content, title, metadata)
        elif format == 'docx':
            return cls._to_docx(content, title, metadata)
    
    @classmethod
    def export_conversation(
        cls,
        messages: List[Dict],
        format: str = 'pdf',
        title: str = "Conversación Zynthra-AI"
    ) -> bytes:
        """Export a full conversation to a document.
        
        Args:
            messages: List of {role, content} dicts
            format: Target format
            title: Document title
            
        Returns:
            Raw bytes of the generated document
        """
        lines = [f"# {title}\n"]
        lines.append(f"*Generado: {datetime.now().strftime('%d/%m/%Y %H:%M')}*\n")
        lines.append("---\n")
        
        for msg in messages:
            role = msg.get('role', 'user')
            content = msg.get('content', '')
            label = "👤 Usuario" if role == 'user' else "🤖 Zynthra-AI"
            lines.append(f"## {label}\n")
            lines.append(f"{content}\n")
            lines.append("---\n")
        
        full_text = "\n".join(lines)
        return cls.export(full_text, format=format, title=title)
    
    @staticmethod
    def _to_txt(content: str, title: Optional[str], metadata: Optional[Dict]) -> bytes:
        """Plain text export."""
        lines = []
        if title:
            lines.append(title)
            lines.append("=" * len(title))
            lines.append("")
        if metadata:
            for k, v in metadata.items():
                lines.append(f"{k}: {v}")
            lines.append("")
        lines.append(content)
        return "\n".join(lines).encode('utf-8')
    
    @staticmethod
    def _to_markdown(content: str, title: Optional[str], metadata: Optional[Dict]) -> bytes:
        """Markdown export."""
        lines = []
        if title:
            lines.append(f"# {title}")
            lines.append("")
        if metadata:
            for k, v in metadata.items():
                lines.append(f"- **{k}:** {v}")
            lines.append("")
        lines.append(content)
        return "\n".join(lines).encode('utf-8')
    
    @staticmethod
    def _to_pdf(content: str, title: Optional[str], metadata: Optional[Dict]) -> bytes:
        """PDF export using reportlab."""
        try:
            from reportlab.lib.pagesizes import A4
            from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
            from reportlab.lib.units import cm
            from reportlab.lib.colors import HexColor
            from reportlab.platypus import (
                SimpleDocTemplate, Paragraph, Spacer, PageBreak
            )
            from reportlab.lib.enums import TA_LEFT, TA_CENTER
        except ImportError:
            logger.error("reportlab not installed. Run: pip install reportlab")
            raise
        
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(
            buffer,
            pagesize=A4,
            rightMargin=2*cm,
            leftMargin=2*cm,
            topMargin=2*cm,
            bottomMargin=2*cm,
            title=title or "Documento Zynthra-AI"
        )
        
        # Estilos Zynthra
        styles = getSampleStyleSheet()
        olive = HexColor('#4A6B2F')
        dark = HexColor('#1F2A1A')
        muted = HexColor('#8A9680')
        
        title_style = ParagraphStyle(
            'ZynthraTitle',
            parent=styles['Heading1'],
            fontSize=20,
            textColor=olive,
            spaceAfter=12,
            alignment=TA_LEFT
        )
        body_style = ParagraphStyle(
            'ZynthraBody',
            parent=styles['BodyText'],
            fontSize=11,
            textColor=dark,
            spaceAfter=10,
            leading=16
        )
        meta_style = ParagraphStyle(
            'ZynthraMeta',
            parent=styles['Italic'],
            fontSize=9,
            textColor=muted,
            spaceAfter=4
        )
        heading_style = ParagraphStyle(
            'ZynthraH2',
            parent=styles['Heading2'],
            fontSize=14,
            textColor=olive,
            spaceAfter=8,
            spaceBefore=12
        )
        
        story = []
        
        if title:
            story.append(Paragraph(title, title_style))
        
        if metadata:
            for k, v in metadata.items():
                story.append(Paragraph(f"<b>{k}:</b> {v}", meta_style))
            story.append(Spacer(1, 0.3*cm))
        
        # Render content line-by-line (basic markdown support)
        for line in content.split('\n'):
            stripped = line.strip()
            if not stripped:
                story.append(Spacer(1, 0.15*cm))
                continue
            
            if stripped.startswith('# '):
                story.append(Paragraph(stripped[2:], title_style))
            elif stripped.startswith('## '):
                story.append(Paragraph(stripped[3:], heading_style))
            elif stripped.startswith('### '):
                story.append(Paragraph(stripped[4:], heading_style))
            elif stripped.startswith('---'):
                story.append(Spacer(1, 0.2*cm))
            elif stripped.startswith('- ') or stripped.startswith('* '):
                story.append(Paragraph(f"• {stripped[2:]}", body_style))
            else:
                # Inline replacements: **bold** → <b>bold</b>
                html = stripped
                # Bold
                while '**' in html:
                    html = html.replace('**', '<b>', 1)
                    if '**' in html:
                        html = html.replace('**', '</b>', 1)
                story.append(Paragraph(html, body_style))
        
        # Footer on each page
        def footer(canvas, doc):
            canvas.saveState()
            canvas.setFont('Helvetica', 8)
            canvas.setFillColor(muted)
            canvas.drawCentredString(
                A4[0] / 2,
                1 * cm,
                f"Generado por Zynthra-AI  •  Página {doc.page}"
            )
            canvas.restoreState()
        
        doc.build(story, onFirstPage=footer, onLaterPages=footer)
        
        pdf_bytes = buffer.getvalue()
        buffer.close()
        return pdf_bytes
    
    @staticmethod
    def _to_docx(content: str, title: Optional[str], metadata: Optional[Dict]) -> bytes:
        """DOCX export using python-docx."""
        try:
            from docx import Document
            from docx.shared import Pt, RGBColor, Cm
            from docx.enum.text import WD_ALIGN_PARAGRAPH
        except ImportError:
            logger.error("python-docx not installed. Run: pip install python-docx")
            raise
        
        doc = Document()
        
        # Page margins
        for section in doc.sections:
            section.top_margin = Cm(2)
            section.bottom_margin = Cm(2)
            section.left_margin = Cm(2.5)
            section.right_margin = Cm(2.5)
        
        olive = RGBColor(0x4A, 0x6B, 0x2F)
        muted = RGBColor(0x8A, 0x96, 0x80)
        
        if title:
            title_p = doc.add_paragraph()
            run = title_p.add_run(title)
            run.font.size = Pt(22)
            run.font.bold = True
            run.font.color.rgb = olive
        
        if metadata:
            for k, v in metadata.items():
                meta_p = doc.add_paragraph()
                run = meta_p.add_run(f"{k}: ")
                run.bold = True
                run.font.size = Pt(9)
                run.font.color.rgb = muted
                run2 = meta_p.add_run(str(v))
                run2.font.size = Pt(9)
                run2.font.color.rgb = muted
            doc.add_paragraph()
        
        # Render content
        for line in content.split('\n'):
            stripped = line.strip()
            if not stripped:
                doc.add_paragraph()
                continue
            
            if stripped.startswith('# '):
                h = doc.add_heading(stripped[2:], level=1)
                for run in h.runs:
                    run.font.color.rgb = olive
            elif stripped.startswith('## '):
                h = doc.add_heading(stripped[3:], level=2)
                for run in h.runs:
                    run.font.color.rgb = olive
            elif stripped.startswith('### '):
                h = doc.add_heading(stripped[4:], level=3)
                for run in h.runs:
                    run.font.color.rgb = olive
            elif stripped.startswith('---'):
                doc.add_paragraph()
            elif stripped.startswith('- ') or stripped.startswith('* '):
                doc.add_paragraph(stripped[2:], style='List Bullet')
            else:
                # Handle inline **bold**
                p = doc.add_paragraph()
                parts = stripped.split('**')
                for i, part in enumerate(parts):
                    run = p.add_run(part)
                    if i % 2 == 1:  # Odd parts are bold
                        run.bold = True
        
        buffer = io.BytesIO()
        doc.save(buffer)
        docx_bytes = buffer.getvalue()
        buffer.close()
        return docx_bytes
    
    @staticmethod
    def get_mime_type(format: str) -> str:
        """Get MIME type for a format."""
        mimes = {
            'txt': 'text/plain',
            'md': 'text/markdown',
            'pdf': 'application/pdf',
            'docx': 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
        }
        return mimes.get(format.lower(), 'application/octet-stream')
